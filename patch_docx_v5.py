# -*- coding: utf-8 -*-
"""在 v4 定稿 docx 上补入作者贡献声明，输出 v5（最终投稿稿）。"""
from docx import Document

SRC = '双重去杠杆陷阱_投稿稿_v4.docx'
DST = '双重去杠杆陷阱_投稿稿_v5.docx'

OLD = '【投稿时填写：多位作者时按《金融研究》体例列明。】'
NEW = '潘郑麒负责论文撰写；潘铁军负责回测代码编写与实证结果验证；郑蕾娜提出研究选题与总体研究思路。'


def main():
    doc = Document(SRC)
    hit = 0
    for p in doc.paragraphs:
        full = ''.join(r.text for r in p.runs)
        if OLD in full:
            p.runs[0].text = full.replace(OLD, NEW)
            for r in p.runs[1:]:
                r.text = ''
            hit += 1
    print('替换命中:', hit)
    doc.save(DST)

    chk = Document(DST)
    for p in chk.paragraphs:
        t = ''.join(r.text for r in p.runs)
        if '潘郑麒' in t:
            print('✓', t[:80])
    rest = 0
    for p in chk.paragraphs:
        t = ''.join(r.text for r in p.runs)
        if '投稿时填写' in t or '［机构待补］' in t:
            rest += 1
    print('残留占位符:', rest, '| 总段落:', len(chk.paragraphs), '| 表格:', len(chk.tables))
    print('生成:', DST)


if __name__ == '__main__':
    main()
