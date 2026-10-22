# -*- coding: utf-8 -*-
"""校验生成的投稿稿 Word: 表格/图/公式数量、文字完整性、章节齐全"""
import os
import re
from docx import Document

BASE = os.path.dirname(os.path.abspath(__file__))
DST = os.path.join(BASE, '双重去杠杆陷阱_投稿稿_v3.docx')
SRC = os.path.join(BASE, '论文_双重去杠杆陷阱.md')

doc = Document(DST)
body = doc.element.body

texts = []
for p in doc.paragraphs:
    texts.append(p.text)
tbl_texts = []
for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            tbl_texts.append(c.text)

all_txt = '\n'.join(texts) + '\n' + '\n'.join(tbl_texts)
cn_doc = len(re.findall(r'[\u4e00-\u9fff]', all_txt))
md = open(SRC, encoding='utf-8').read()
cn_md = len(re.findall(r'[\u4e00-\u9fff]', md))

print('表格数         :', len(doc.tables), '  (期望 24)')
print('图片(内联)     :', len(body.findall(
    './/{http://schemas.openxmlformats.org/drawingml/2006/main}blip')))
print('公式 oMathPara :', len(body.findall(
    './/{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara')))
print('段落数         :', len(doc.paragraphs))
print('中文字符 docx  :', cn_doc, '  md:', cn_md, '  差:', cn_doc - cn_md)

must = ['摘　要', '一、引言', '1.5 边际贡献', '2.3 中国市场的趋势与动量',
        '2.4 理论模型', '四、实证结果', '5.4 回撤控制成本的干净分离', '5.5 数值模拟', '6.7 参数与成本敏感性',
        '7.5 回撤控制何以', '8.2 理论贡献',
        '6.6 关闭战术信号', '七、扩展讨论：约束成本的量化',
        '八、结论与启示', '8.4 研究局限',
        '参考文献', '图1　', '图2　', '图3　', '图4　']
miss = [m for m in must if m not in all_txt]
print('章节缺失       :', miss if miss else '无 ✓')

caps = re.findall(r'表\s?\d+　', all_txt)
print('表题数         :', len(caps))

# 关键数字是否有丢失
nums = ['4.21亿户', '3.54万亿', '−0.6209', '0.821', '0.461', '+0.439',
        '3.794', '0.363', '1.12pp', '5.05pp', '269亿元', '2.475']
lost = [n for n in nums if n not in all_txt]
print('关键数字缺失   :', lost if lost else '无 ✓')

# 异常长段落（可能吞并了公式/表格）
long_p = [t[:60] for t in texts if len(t) > 900]
print('超长段落(>900) :', len(long_p))
for t in long_p[:3]:
    print('   ', t)
