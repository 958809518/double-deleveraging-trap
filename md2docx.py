# -*- coding: utf-8 -*-
"""
把论文 Markdown 转为《金融研究》投稿体例的 Word 文档
- 正文: 宋体/Times New Roman 小四(12pt), 1.5倍行距, 首行缩进2字符
- 标题: 一级黑体15pt居中 / 二级黑体13pt / 三级黑体12pt
- 表格: 三线表(顶线+表头下线+底线), 五号宋体, 表题居上
- 图:   居中嵌入 300dpi PNG, 图题居下
- 公式: LaTeX -> MathML -> OMML, 原生 Word 公式对象
- 参考文献: 悬挂缩进
"""
import os
import re
import latex2mathml.converter as L2M
from mathml2omml import convert as M2O
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, '论文_双重去杠杆陷阱.md')
DST = os.path.join(BASE, '双重去杠杆陷阱_投稿稿_v4.docx')
FIGDIR = os.path.join(BASE, 'paper_data')

NS_M = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
BODY_CN, BODY_EN = '宋体', 'Times New Roman'
HEAD_CN = '黑体'
TEXTW = 14.66  # cm 版心宽

FIGCAP = {
    'fig1_frontier.svg': '图1　两种约束集下的有效前沿（圈注：单峰顶点）',
    'fig2_constraint_cost.svg': '图2　约束成本的量化：回撤控制“消灭”了杠杆的价值',
    'fig3_nosignal.svg': '图3　关闭全部战术信号后，前沿分岔依然存在（圈注：单峰顶点）',
    'fig4_boundary.svg': '图4　回撤控制成本的分离与条件性（A：表10；B：表19）',
}


# ---------------- 基础工具 ----------------
def set_run(r, cn=BODY_CN, en=BODY_EN, size=12, bold=False, italic=False):
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    rPr = r._element.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    rf.set(qn('w:ascii'), en)
    rf.set(qn('w:hAnsi'), en)
    rf.set(qn('w:eastAsia'), cn)


def para(doc, align=None, before=0, after=0, indent=0, line=1.5, hanging=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if hanging is not None:
        pf.left_indent = Pt(hanging)
        pf.first_line_indent = Pt(-hanging)
    elif indent:
        pf.first_line_indent = Pt(indent)
    return p


TOKEN = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+?\*)')


def rich(p, text, size=12, cn=BODY_CN, bold=False, ital=False):
    """解析 **加粗** / `代码` / *斜体*(仅当 ital=True，避免误伤 σ*)"""
    for tk in TOKEN.split(text):
        if not tk:
            continue
        b, it = bold, False
        if tk.startswith('**') and tk.endswith('**') and len(tk) > 4:
            tk, b = tk[2:-2], True
        elif tk.startswith('`') and tk.endswith('`') and len(tk) > 2:
            tk = tk[1:-1]
        elif ital and tk.startswith('*') and tk.endswith('*') and len(tk) > 2:
            tk, it = tk[1:-1], True
        set_run(p.add_run(tk), cn=cn, size=size, bold=b, italic=it)
    return p


def cell_border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        if edge not in edges:
            continue
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        for k, v in edges[edge].items():
            el.set(qn('w:' + k), str(v))
        if edges[edge] is None:
            borders.remove(el)


def no_border(cell):
    for e in ('top', 'left', 'bottom', 'right'):
        cell_border(cell, **{e: {'val': 'nil'}})


# mathml2omml 已知 bug: <m:groupChrPr>...</m:groupChr> 未正确闭合
FIX_GROUPCHR = re.compile(
    r'(<m:groupChrPr>(?:<m:chr[^>]*/>)?\s*(?:<m:pos[^>]*/>)?)</m:groupChr>')


def add_eq(p, tex):
    inner = M2O(L2M.convert(tex)).strip()
    inner = re.sub(r'^<m:oMath>', '', inner)
    inner = re.sub(r'</m:oMath>$', '', inner)
    inner = FIX_GROUPCHR.sub(r'\1</m:groupChrPr>', inner)
    xml = f'<m:oMathPara {NS_M}>{inner}</m:oMathPara>'
    p._p.append(parse_xml(xml))


def add_field(p, instr, size=10.5):
    r = p.add_run()
    f = OxmlElement('w:fldChar')
    f.set(qn('w:fldCharType'), 'begin')
    r._r.append(f)
    r2 = p.add_run()
    t = OxmlElement('w:instrText')
    t.set(qn('xml:space'), 'preserve')
    t.text = instr
    r2._r.append(t)
    set_run(r2, size=size)
    r3 = p.add_run()
    f2 = OxmlElement('w:fldChar')
    f2.set(qn('w:fldCharType'), 'end')
    r3._r.append(f2)


# ---------------- 文档骨架 ----------------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.54)
sec.left_margin = sec.right_margin = Cm(3.17)

st = doc.styles['Normal']
st.font.size = Pt(12)
st.font.name = BODY_EN
st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_CN)

# 页眉页脚
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(hp.add_run('双重去杠杆陷阱：波动率目标与回撤控制的交互效应'), cn=BODY_CN, size=9)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_field(fp, ' PAGE ', size=10.5)

# 首页信息
p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=2, line=1.25)
set_run(p.add_run('双重去杠杆陷阱：波动率目标与回撤控制的交互效应'),
        cn=HEAD_CN, size=18, bold=True)
p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=10, line=1.25)
set_run(p.add_run('——来自中国多资产市场的证据与数值模拟'), cn='楷体', size=13)

p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=2, line=1.25)
set_run(p.add_run('作者姓名A　　作者姓名B'), cn='楷体', size=11)
p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=14, line=1.25)
set_run(p.add_run('（1. 作者单位，城市 邮编；2. 作者单位，城市 邮编）'),
        cn='楷体', size=9.5)
note = para(doc, before=0, after=14, line=1.25)
set_run(note.add_run(
    '【投稿前请填写作者与单位信息；若为匿名审稿稿，请删除上述四行】'),
    cn='楷体', size=9, italic=True)

# ---------------- 解析 ----------------
lines = open(SRC, encoding='utf-8').read().split('\n')
i = 0
in_abstract = False
in_ref = False

while i < len(lines):
    ln = lines[i].rstrip()
    s = ln.strip()

    # 顶部体例说明块 -> 跳过
    if s.startswith('> 体例') or s.startswith('> 数据') or s.startswith('> 所有回测'):
        while i < len(lines) and lines[i].strip().startswith('>'):
            i += 1
        continue

    if s in ('---', ''):
        i += 1
        continue

    # 主标题 / 副标题（已写入首页，跳过）
    if ln.startswith('# ') or ln.startswith('## ——'):
        i += 1
        continue

    # 一级标题
    if ln.startswith('## '):
        t = ln[3:].strip()
        in_ref = (t == '参考文献')
        in_abstract = (t == '摘要')
        if t == '摘要':
            p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6, line=1.25)
            set_run(p.add_run('摘　要'), cn=HEAD_CN, size=12, bold=True)
        elif t == '参考文献':
            p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=16, after=8, line=1.25)
            set_run(p.add_run('参考文献'), cn=HEAD_CN, size=15, bold=True)
        else:
            p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=16, after=10)
            set_run(p.add_run(t), cn=HEAD_CN, size=15, bold=True)
        i += 1
        continue

    # 二级标题
    if ln.startswith('### '):
        p = para(doc, before=12, after=5, indent=0)
        set_run(p.add_run(ln[4:].strip()), cn=HEAD_CN, size=13, bold=True)
        i += 1
        continue

    # 三级标题
    if ln.startswith('#### '):
        p = para(doc, before=9, after=4, indent=0)
        set_run(p.add_run(ln[5:].strip()), cn=HEAD_CN, size=12, bold=True)
        i += 1
        continue

    # 配图
    if ln.startswith('@@FIG'):
        rel = ln[5:].strip()
        name = os.path.basename(rel)
        fname = {'fig1_frontier.svg': 'docx_fig1.png',
                 'fig2_constraint_cost.svg': 'docx_fig2.png',
                 'fig3_nosignal.svg': 'docx_fig3.png',
                 'fig4_boundary.svg': 'docx_fig4.png'}.get(name)
        path = os.path.join(FIGDIR, fname or name)
        w = Cm(11.0) if 'fig3' in (fname or '') else Cm(TEXTW)
        if os.path.exists(path):
            p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=10, after=3, line=1.0)
            p.add_run().add_picture(path, width=w)
            cap = FIGCAP.get(name)
            if cap:
                cp = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=2, after=10,
                          line=1.15, indent=0)
                set_run(cp.add_run(cap), cn=HEAD_CN, size=10.5, bold=True)
        i += 1
        continue

    # 块级公式
    if s.startswith('$$'):
        if len(s) > 4 and s.endswith('$$'):
            tex = s.strip('$').strip()
            i += 1
        else:
            buf = []
            while i < len(lines):
                buf.append(lines[i].strip())
                if len(buf) > 1 and lines[i].strip().endswith('$$'):
                    break
                i += 1
            tex = '\n'.join(buf).strip().strip('$').strip()
            i += 1
        p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=8, line=1.0,
                 indent=0)
        try:
            add_eq(p, tex)
        except Exception as e:
            set_run(p.add_run(tex), size=11)
        continue

    # 表题 + 表格（题与表之间可能空行）
    if s.startswith('**表'):
        j = i + 1
        while j < len(lines) and lines[j].strip() == '':
            j += 1
        if not (j < len(lines) and lines[j].strip().startswith('|')):
            i += 1
            continue
        cap = s.strip('*')
        p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=10, after=3, line=1.15,
                 indent=0)
        set_run(p.add_run(cap), cn=HEAD_CN, size=10.5, bold=True)
        i = j
        header = [c.strip() for c in lines[i].strip().strip('|').split('|')]
        i += 2
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
            i += 1
        ncol = len(header)
        fsize = 8.5 if ncol >= 7 else (9.5 if ncol >= 5 else 10.5)
        tbl = doc.add_table(rows=1 + len(rows), cols=ncol)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True
        tblW = OxmlElement('w:tblW')
        tblW.set(qn('w:w'), '5000')
        tblW.set(qn('w:type'), 'pct')
        tbl._tbl.tblPr.append(tblW)
        for j, h in enumerate(header):
            c = tbl.cell(0, j)
            c.text = ''
            cp = c.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.line_spacing = 1.0
            cp.paragraph_format.space_before = Pt(1.5)
            cp.paragraph_format.space_after = Pt(1.5)
            rich(cp, re.sub(r'\*\*|\*', '', h), size=fsize, bold=True)
            no_border(c)
            cell_border(c, top={'val': 'single', 'sz': 8, 'color': '000000'},
                        bottom={'val': 'single', 'sz': 4, 'color': '000000'})
        for ri, row in enumerate(rows):
            for j in range(ncol):
                c = tbl.cell(ri + 1, j)
                c.text = ''
                cp = c.paragraphs[0]
                cp.alignment = (WD_ALIGN_PARAGRAPH.LEFT if j == 0
                                else WD_ALIGN_PARAGRAPH.CENTER)
                cp.paragraph_format.line_spacing = 1.0
                cp.paragraph_format.space_before = Pt(1)
                cp.paragraph_format.space_after = Pt(1)
                rich(cp, re.sub(r'\*\*|\*', '', row[j] if j < len(row) else ''),
                     size=fsize)
                no_border(c)
                if ri == len(rows) - 1:
                    cell_border(c, bottom={'val': 'single', 'sz': 8,
                                           'color': '000000'})
        continue

    # 图题（紧跟图片之后的独立粗体行）
    if s.startswith('**图'):
        p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=2, after=10, line=1.15,
                 indent=0)
        set_run(p.add_run(s.strip('*')), cn=HEAD_CN, size=10.5, bold=True)
        i += 1
        continue

    # 引用块
    if ln.startswith('>'):
        p = para(doc, before=6, after=6, indent=0, line=1.4)
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.right_indent = Pt(24)
        set_run(p.add_run(ln.lstrip('> ').strip()), cn='楷体', size=11)
        i += 1
        continue

    # 列表项
    if re.match(r'^[-*]\s+', ln):
        p = para(doc, before=2, after=2, line=1.4, indent=0)
        p.paragraph_format.left_indent = Pt(24)
        rich(p, re.sub(r'^[-*]\s+', '', ln), size=12)
        i += 1
        continue

    # 参考文献标题（原文为加粗段落而非标题语法）
    if s == '**参考文献**':
        in_ref = True
        p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=16, after=8, line=1.25)
        set_run(p.add_run('参考文献'), cn=HEAD_CN, size=15, bold=True)
        i += 1
        continue

    # 参考文献条目
    if in_ref and re.match(r'^\d+\.\s', s):
        p = para(doc, before=0, after=2, line=1.3, hanging=24, indent=0)
        rich(p, s, size=10.5, ital=True)
        i += 1
        continue

    # 关键词 / JEL
    if s.startswith('**关键词**') or s.startswith('**JEL'):
        p = para(doc, before=4, after=4, line=1.4, indent=0)
        rich(p, s, size=10.5)
        i += 1
        continue

    # 正文段落
    if in_abstract:
        p = para(doc, before=0, after=0, line=1.4, indent=21)
        p.paragraph_format.left_indent = Pt(0)
        rich(p, s, size=10.5)
    else:
        p = para(doc, before=0, after=3, line=1.5, indent=24)
        rich(p, s, size=12)
    i += 1

# ---------- 英文标题 / 摘要 / 关键词（《金融研究》体例，置于文末）----------
EN_TITLE = ('Double Deleveraging Trap: The Interaction between Volatility '
            'Targeting and Drawdown Control')
EN_SUB = "Evidence from China's Multi-Asset Market"
EN_ABS = (
    'Volatility targeting and drawdown control are the two most widely used '
    'rule-based risk overlays in asset management, yet prior studies typically '
    'evaluate them in isolation and almost universally assume that investors can '
    'employ leverage. Building on the debate between Moreira and Muir (2017) and '
    'Cederburg et al. (2020), this paper places both overlays within a unified '
    'no-leverage framework and examines their interaction. Theory shows that '
    'because the position-scaling function of volatility targeting is truncated '
    'at the leverage cap, its "leverage up in calm markets" channel is disabled '
    'and only the deleveraging channel survives; because drawdown control cuts '
    'exposure multiplicatively, the two layers compound into a "double '
    'deleveraging". Using a 13.2-year sample of six Chinese asset classes plus '
    'treasury and cash, we find that with volatility targeting alone the '
    'return-risk frontier is monotonically increasing (target volatility 3%-12% '
    'corresponds to annualized returns of 5.48%-11.28%), whereas adding drawdown '
    'control turns the frontier single-peaked: the interior optimum occurs at a '
    'target volatility of 5.5% (annualized 6.47%, maximum drawdown -4.15%), and '
    'raising the target volatility further reduces rather than increases returns '
    '(to 5.37% at 12%). A descriptive slope decomposition shows that drawdown '
    'control shifts the return-target-volatility slope by -0.6209, offsetting '
    '95.8% of the main effect (+0.6479). Mechanism tests show that drawdown '
    'control lowers the exposure realization rate from 0.821 to 0.461: a 92% '
    'increase in target exposure buys only a 6.9% increase in effective '
    'exposure. An event study with state controls and block bootstrap indicates '
    'that the loss does not arise from mistimed cuts -- drawdown control reduces '
    'exposure at moments when subsequent 20-120-day market returns are in fact '
    'lower -- but from the compounding loss and ratchet effect of compressed '
    'exposure. Quantifying the welfare cost of the two constraints under three '
    'metrics yields a counterintuitive result: once a drawdown constraint is '
    'imposed, the welfare cost of the leverage constraint falls to nearly zero, '
    'while the cost of drawdown control itself reaches 3.794 percentage points, '
    'more than ten times the former. Two caveats are stated up front: the '
    "strategy's risk-adjusted advantage over a treasury ETF is not statistically "
    'significant (p = 0.774), and the single-peaked shape is conditional -- it '
    'appears under two of four drawdown-threshold settings and five of six '
    'estimation, frequency and functional-form variants, although the cost of '
    'drawdown control at high target volatility (4-7 percentage points) holds '
    'without exception. Our findings imply that, absent leverage, risk budgets '
    'are not "the more the better", and stacking additional risk-control layers '
    'systematically erodes returns, with direct implications for retail investor '
    'protection and the design of "fixed income+" products.')
EN_KEY = ('volatility targeting; drawdown control; portfolio insurance; '
          'asset allocation; retail investor constraints')

p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, before=14, after=0, line=1.3, indent=0)
set_run(p.add_run('说明：本文所有回测均已扣除交易成本（单边万三），并校正了 ETF '
                  '份额折算造成的价格假跳变；数据区间 2013-08-01 至 2026-10-20。'),
        cn='楷体', size=9.5)

p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=18, after=2, line=1.25, indent=0)
set_run(p.add_run(EN_TITLE), size=13, bold=True)
p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=0, after=8, line=1.25, indent=0)
set_run(p.add_run('——' + EN_SUB), size=11)

p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=4, line=1.25, indent=0)
set_run(p.add_run('Abstract: '), size=10.5, bold=True)
set_run(p.add_run(EN_ABS), size=10.5)

p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=3, line=1.25, indent=0)
set_run(p.add_run('Keywords: '), size=10.5, bold=True)
set_run(p.add_run(EN_KEY), size=10.5)

p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, line=1.25, indent=0)
set_run(p.add_run('JEL Classification: '), size=10.5, bold=True)
set_run(p.add_run('G11, G12, G23'), size=10.5)

doc.save(DST)
print('生成:', DST, os.path.getsize(DST), 'bytes')
