# -*- coding: utf-8 -*-
"""在 v3 定稿 docx 上定点补填投稿信息，输出 v4。
避免重新走 md2docx（其依赖 latex2mathml，当前环境缺失且无法联网安装）。
"""
from docx import Document

SRC = '双重去杠杆陷阱_投稿稿_v3.docx'
DST = '双重去杠杆陷阱_投稿稿_v4.docx'

REPLACEMENTS = [
    # 正文引用年份与文献表同步
    ('Xu（2024）则给出实时改进的缩放方案。',
     'Xu（2026）则给出实时改进的缩放方案。'),
    # 文献 7：作者名缩写、标题、出处、SSRN 编号修正
    ('7. Liu, J., Tang, Y., & Zhou, G. (2019). Volatility-managed portfolios: '
     'Does it scale? Working paper.［机构待补］',
     '7. Liu, F., Tang, X., & Zhou, G. (2019). Volatility-managed portfolio: '
     'Does it really work? Journal of Portfolio Management, 46(1), 38–51.'
     '（SSRN 工作论文版编号 3283395）'),
    # 文献 11：年份、出版出处、SSRN 编号修正
    ('11. Xu, X. (2024). Improving volatility-managed portfolios in real time. '
     'Working paper.［机构待补］',
     '11. Xu, X. (2026). Improving volatility-managed portfolios in real time. '
     'Critical Finance Review, 15(2), 179–207.（SSRN 工作论文版编号 4778941）'),
    # 上交所数据访问日期
    ('上交所官网公开数据（访问日期投稿时补）。',
     '上交所官网公开数据（访问日期：2026年10月22日）。'),
    # 数据可用性：填入代码仓库地址
    ('回测与绘图代码已在作者主页公开（投稿时填写：代码仓库地址），可复现文中所有表格与图形。',
     '回测与绘图代码已在 GitHub 公开：https://github.com/958809518/double-deleveraging-trap，'
     '含 paper_data/ 下与正文各表一一对应的结果输出，可复现文中所有表格与图形。'),
    # 资助声明
    ('【投稿时填写：基金项目名称与批准号。】',
     '本文系国家社会科学基金一般项目《数字孪生赋能韧性城市建设的机理与实践路径研究》'
     '（批准号 24BGL226）的阶段性成果。'),
]


def main():
    doc = Document(SRC)
    done = set()
    for p in doc.paragraphs:
        full = ''.join(r.text for r in p.runs)
        if not full:
            continue
        newfull = full
        for old, new in REPLACEMENTS:
            if old in newfull:
                newfull = newfull.replace(old, new)
                done.add(old)
        if newfull != full and p.runs:
            p.runs[0].text = newfull
            for r in p.runs[1:]:
                r.text = ''

    print('已完成替换: %d/%d' % (len(done), len(REPLACEMENTS)))
    for old, _ in REPLACEMENTS:
        if old not in done:
            print('  [未命中]', old[:50])

    doc.save(DST)
    print('生成:', DST)

    # 复核
    chk = Document(DST)
    for p in chk.paragraphs:
        t = ''.join(r.text for r in p.runs)
        for k in ['Liu, F.', 'Xu, X. (2026)', '访问日期：2026年10月22日',
                  'double-deleveraging-trap', '24BGL226', 'Xu（2026）']:
            if k in t:
                print('  ✓', k)
                break


if __name__ == '__main__':
    main()
